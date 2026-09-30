"""Build the live-formula Excel model: py -3 build_workbook.py

Every calculation cell is an Excel formula. Only the Inputs sheet holds typed numbers.
Inputs here mirror python/waterfall/config.py; python/tests checks that they agree.
"""
import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

# ============================== CONFIG ==============================
OUTPUT_PATH = Path(__file__).resolve().parent / "securitization_waterfall.xlsx"
AUTHOR = "Yazad Madan"
SOFFICE_CANDIDATES = (r"C:\Program Files\LibreOffice\program\soffice.exe",
                      r"C:\Program Files (x86)\LibreOffice\program\soffice.exe", "soffice")
HORIZON = 96
ROW0 = 5  # sheet row holding period 0 on every time-series sheet

REP_LINES = [  # name, balance USD, WAC, remaining term months (SYNTHETIC)
    ("RL01", 60_000_000, 0.0650, 60), ("RL02", 55_000_000, 0.0725, 72), ("RL03", 50_000_000, 0.0800, 48),
    ("RL04", 48_000_000, 0.0875, 66), ("RL05", 45_000_000, 0.0950, 60), ("RL06", 40_000_000, 0.1025, 84),
    ("RL07", 35_000_000, 0.1100, 54), ("RL08", 30_000_000, 0.1175, 72), ("RL09", 22_000_000, 0.1250, 36),
    ("RL10", 15_000_000, 0.1350, 48),
]
SCENARIO_TABLE = [  # name, CPR, CDR, severity, recovery lag months, forward shift bps
    ("Base", 0.10, 0.03, 0.40, 6, 0), ("Mild stress", 0.10, 0.06, 0.50, 6, 0),
    ("Stress", 0.06, 0.12, 0.60, 9, 0), ("Severe stress", 0.04, 0.20, 0.70, 12, 0),
    ("Fast prepay", 0.30, 0.03, 0.40, 6, 0), ("Rate shock +300bp", 0.10, 0.06, 0.50, 6, 300),
]
SELECTED_SCENARIO = 1  # 1-based row of SCENARIO_TABLE
SCALARS = {
    "fwd_start": 0.0430, "fwd_end": 0.0325, "fwd_tau": 24.0,
    "bal_a": 300_000_000, "bal_b": 40_000_000, "bal_c": 30_000_000,
    "margin_a": 0.0110, "margin_b": 0.0225, "fixed_c": 0.0750,
    "svc": 0.0060, "trustee": 0.0003, "res_init": 3_000_000, "res_target": 3_000_000,
    "oc_a": 1.25, "oc_b": 1.12, "oc_c": 1.045, "ic_a": 1.50, "ic_b": 1.30, "ic_c": 1.15,
    "fvm_a": 0.0110, "fvm_b": 0.0275, "fvm_c": 0.0450, "fvm_r": 0.1200,
    "ret_pct": 0.05, "lshape_v": 0.025,
    "eu_sold_a": 1, "eu_sold_b": 1, "eu_sold_c": 1, "eu_sold_r": 1,
    "adv_a": 0.85, "adv_b": 0.70, "adv_c": 0.55, "adv_r": 0.30, "fin_spread": 0.0275,
}
PRINCIPAL_MODE = "sequential"
# ============================ END CONFIG ============================

TR = ("a", "b", "c")
ALL = ("a", "b", "c", "r")
BOLD = Font(bold=True)
NUM = "#,##0;(#,##0);-"
PCT = "0.0000%"
RATIO = "0.0000"


class Inputs:
    """Inputs sheet writer that remembers each named cell's absolute address."""

    def __init__(self, ws):
        self.ws, self.row, self.addr = ws, 1, {}

    def title(self, text):
        self.ws.cell(self.row, 1, text).font = BOLD
        self.row += 1

    def put(self, name, label, value, note="", fmt=None):
        self.ws.cell(self.row, 1, label)
        c = self.ws.cell(self.row, 2, value)
        if fmt:
            c.number_format = fmt
        self.ws.cell(self.row, 3, note)
        self.addr[name] = f"Inputs!$B${self.row}"
        self.row += 1

    def blank(self):
        self.row += 1


class Series:
    """A time-series sheet: columns are added by name, formulas are lambdas of the row number."""

    def __init__(self, wb, title, note):
        self.ws = wb.create_sheet(title)
        self.title = title
        self.cols, self.specs = {}, []
        self.ws.cell(1, 1, title).font = BOLD
        self.ws.cell(2, 1, note)

    def add(self, name, header, row0, rowt, fmt=NUM):
        idx = len(self.cols) + 1
        self.cols[name] = get_column_letter(idx)
        self.specs.append((idx, header, row0, rowt, fmt))

    def a(self, name, r):
        return f"{self.cols[name]}{r}"

    def x(self, name, r):
        """Cross-sheet reference."""
        return f"{self.title}!{self.cols[name]}{r}"

    def rng(self, name, first=ROW0, last=ROW0 + HORIZON, sheet=True):
        col = self.cols[name]
        body = f"${col}${first}:${col}${last}"
        return f"{self.title}!{body}" if sheet else body

    def fill(self):
        for idx, header, row0, rowt, fmt in self.specs:
            h = self.ws.cell(4, idx, header)
            h.font = BOLD
            h.alignment = Alignment(wrap_text=True)
            self.ws.column_dimensions[get_column_letter(idx)].width = 16
            for t in range(HORIZON + 1):
                r = ROW0 + t
                src = row0 if t == 0 else rowt
                v = src(r) if callable(src) else src
                if v is None:
                    continue
                c = self.ws.cell(r, idx, v)
                c.number_format = fmt
        self.ws.freeze_panes = self.ws.cell(ROW0, 2)


def irr12(rng):
    """Annualized monthly IRR with fallback guesses so deep-loss tranches still solve."""
    return (f'=IFERROR(12*IRR({rng},0.005),IFERROR(12*IRR({rng},-0.03),IFERROR(12*IRR({rng},-0.1),IFERROR(12*IRR({rng},-0.3),"n/a"))))')


def build(path=OUTPUT_PATH):
    wb = Workbook()
    wb.properties.creator = AUTHOR
    wb.properties.title = "Securitization waterfall and risk retention model"
    wb.calculation.fullCalcOnLoad = True
    ws_in = wb.active
    ws_in.title = "Inputs"
    inp = Inputs(ws_in)
    I = inp.addr  # noqa: E741

    # Inputs
    inp.title("Inputs (typed values live only on this sheet). Collateral is SYNTHETIC; forward curve is ASSUMED.")
    inp.blank()
    inp.title("Active scenario")
    inp.put("scen_sel", "Selected scenario number (1-based row of scenario table)", SELECTED_SCENARIO)

    inp.title("Pool rep lines (name, balance, WAC, remaining term)")
    rep_rows = []
    for name, bal, wac, wam in REP_LINES:
        ws_in.cell(inp.row, 1, name)
        ws_in.cell(inp.row, 2, bal).number_format = NUM
        ws_in.cell(inp.row, 3, wac).number_format = PCT
        ws_in.cell(inp.row, 4, wam)
        rep_rows.append(inp.row)
        inp.row += 1
    inp.blank()
    inp.title("Scenario table (name, CPR, CDR, severity, recovery lag months, forward shift bps)")
    scen_rows = []
    for row in SCENARIO_TABLE:
        for j, v in enumerate(row):
            c = ws_in.cell(inp.row, 1 + j, v)
            if j in (1, 2, 3):
                c.number_format = "0.00%"
        scen_rows.append(inp.row)
        inp.row += 1
    s1, s2 = scen_rows[0], scen_rows[-1]
    inp.blank()
    inp.title("Active scenario values (formulas)")
    sel = I["scen_sel"]
    inp.put("scen_name", "Scenario name", f"=INDEX($A${s1}:$A${s2},{sel})")
    inp.put("cpr", "CPR", f"=INDEX($B${s1}:$B${s2},{sel})", "annual", "0.00%")
    inp.put("cdr", "CDR", f"=INDEX($C${s1}:$C${s2},{sel})", "annual", "0.00%")
    inp.put("sev", "Loss severity", f"=INDEX($D${s1}:$D${s2},{sel})", "of defaulted balance", "0.00%")
    inp.put("lag", "Recovery lag (months)", f"=INDEX($E${s1}:$E${s2},{sel})")
    inp.put("shift", "Forward curve shift (bps)", f"=INDEX($F${s1}:$F${s2},{sel})")
    inp.put("smm", "SMM", f"=1-(1-{I['cpr']})^(1/12)", "monthly", PCT)
    inp.put("mdr", "MDR", f"=1-(1-{I['cdr']})^(1/12)", "monthly", PCT)
    inp.blank()
    inp.title("Curve, notes, fees, reserve")
    labels = {
        "fwd_start": "Forward index at t=0", "fwd_end": "Forward index long-run", "fwd_tau": "Curve time constant (months)",
        "bal_a": "Class A balance", "bal_b": "Class B balance", "bal_c": "Class C balance",
        "margin_a": "Class A margin over index", "margin_b": "Class B margin over index", "fixed_c": "Class C fixed coupon",
        "svc": "Servicing fee (annual)", "trustee": "Trustee fee (annual)",
        "res_init": "Reserve account initial funding", "res_target": "Reserve account target",
    }
    for k, lab in labels.items():
        inp.put(k, lab, SCALARS[k], fmt=NUM if k.startswith(("bal", "res")) else PCT if k != "fwd_tau" else None)
    inp.put("pool_bal", "Pool balance", f"=SUM($B${rep_rows[0]}:$B${rep_rows[-1]})", fmt=NUM)
    inp.put("bal_r", "Residual balance (pool + reserve - A - B - C)",
            f"={I['pool_bal']}+{I['res_init']}-{I['bal_a']}-{I['bal_b']}-{I['bal_c']}", fmt=NUM)
    inp.put("fee_rate", "Senior fee rate (annual)", f"={I['svc']}+{I['trustee']}", fmt=PCT)
    inp.put("mode", "Principal mode (sequential or pro_rata)", PRINCIPAL_MODE)
    inp.put("mode_seq", "Sequential flag", f'=IF({I["mode"]}="sequential",1,0)')
    inp.blank()
    inp.title("Credit enhancement triggers")
    for k, lab in (("oc_a", "OC test Class A minimum"), ("oc_b", "OC test Class B minimum"), ("oc_c", "OC test Class C minimum"),
                   ("ic_a", "IC test Class A minimum"), ("ic_b", "IC test Class B minimum"), ("ic_c", "IC test Class C minimum")):
        inp.put(k, lab, SCALARS[k], fmt=RATIO)
    inp.blank()
    inp.title("Retention, fair value margins, financing")
    for k, lab in (("fvm_a", "FV margin A over forward"), ("fvm_b", "FV margin B"), ("fvm_c", "FV margin C"),
                   ("fvm_r", "FV margin residual"), ("ret_pct", "Required retention"),
                   ("lshape_v", "L-shape vertical share"), ("eu_sold_a", "EU vertical: A sold (1/0)"),
                   ("eu_sold_b", "EU vertical: B sold (1/0)"), ("eu_sold_c", "EU vertical: C sold (1/0)"),
                   ("eu_sold_r", "EU vertical: R sold (1/0)"), ("adv_a", "Advance rate A"), ("adv_b", "Advance rate B"),
                   ("adv_c", "Advance rate C"), ("adv_r", "Advance rate residual"), ("fin_spread", "Financing spread over forward")):
        inp.put(k, lab, SCALARS[k], fmt=None if k.startswith("eu_sold") else PCT)
    inp.blank()
    inp.title("Optional FV sizing override: leave blank to size retention on the selected scenario, or paste base-case fair values to fix sizing at closing")
    for k in ALL:
        inp.put(f"fv_lock_{k}", f"Locked fair value {k.upper()}", None, fmt=NUM)
    ws_in.column_dimensions["A"].width = 58
    ws_in.column_dimensions["B"].width = 18
    ws_in.column_dimensions["C"].width = 14
    ws_in.column_dimensions["D"].width = 14
    ws_in.column_dimensions["E"].width = 14
    ws_in.column_dimensions["F"].width = 14

    per = lambda r: r - ROW0  # noqa: E731

    # Pool
    P = Series(wb, "Pool", "Rep-line amortization. Order: defaults, interest, scheduled principal, prepayment. Period 0 is closing.")
    P.add("per", "Period", lambda r: per(r), lambda r: per(r), "0")
    P.add("fwd", "Forward index", lambda r: f"={I['fwd_start']}+({I['fwd_end']}-{I['fwd_start']})*(1-EXP(-{P.a('per', r)}/{I['fwd_tau']}))+{I['shift']}/10000",
          lambda r: f"={I['fwd_start']}+({I['fwd_end']}-{I['fwd_start']})*(1-EXP(-{P.a('per', r)}/{I['fwd_tau']}))+{I['shift']}/10000", PCT)
    for i, ir in enumerate(rep_rows, start=1):
        bal, wac, wam = f"Inputs!$B${ir}", f"Inputs!$C${ir}", f"Inputs!$D${ir}"
        b, d, n_, s, p_, e = (f"{k}{i}" for k in ("beg", "def", "int", "sch", "pre", "end"))
        P.add(b, f"RL{i:02d} beg", None, lambda r, e=e: f"={P.a(e, r - 1)}")
        P.add(d, f"RL{i:02d} defaults", None, lambda r, b=b: f"={P.a(b, r)}*{I['mdr']}")
        P.add(n_, f"RL{i:02d} interest", None, lambda r, b=b, d=d, wac=wac: f"=({P.a(b, r)}-{P.a(d, r)})*{wac}/12")
        P.add(s, f"RL{i:02d} sched prin", None,
              lambda r, b=b, d=d, wac=wac, wam=wam: (
                  f"=IF(MAX({wam}-({P.a('per', r)}-1),0)<=0,0,({P.a(b, r)}-{P.a(d, r)})*(({wac}/12)/(1-(1+{wac}/12)^(-MAX({wam}-({P.a('per', r)}-1),0)))-{wac}/12))"))
        P.add(p_, f"RL{i:02d} prepay", None,
              lambda r, b=b, d=d, s=s: f"=({P.a(b, r)}-{P.a(d, r)}-{P.a(s, r)})*{I['smm']}")
        P.add(e, f"RL{i:02d} end", f"={bal}",
              lambda r, b=b, d=d, s=s, p_=p_: f"={P.a(b, r)}-{P.a(d, r)}-{P.a(s, r)}-{P.a(p_, r)}")
    nl = len(rep_rows)

    def tot(prefix):
        return lambda r: "=" + "+".join(P.a(f"{prefix}{i}", r) for i in range(1, nl + 1))

    for key, pre, hdr in (("beg", "beg", "Pool beg balance"), ("def", "def", "Pool defaults"), ("int", "int", "Pool interest"),
                          ("sch", "sch", "Pool sched principal"), ("pre", "pre", "Pool prepayments")):
        P.add(f"t_{key}", hdr, None, tot(pre))
    P.add("t_end", "Pool end balance", tot("end"), tot("end"))
    P.add("rec", "Recoveries", None,
          lambda r: f"=IF({P.a('per', r)}-{I['lag']}>=1,INDEX({P.rng('t_def', sheet=False)},{P.a('per', r)}-{I['lag']}+1)*(1-{I['sev']}),0)")
    P.add("cumdef", "Cum defaults", 0, lambda r: f"={P.a('cumdef', r - 1)}+{P.a('t_def', r)}")
    P.add("cumrec", "Cum recoveries", 0, lambda r: f"={P.a('cumrec', r - 1)}+{P.a('rec', r)}")
    P.add("cumloss", "Cum net loss (default x severity)", 0, lambda r: f"={P.a('cumdef', r)}*{I['sev']}")
    P.add("pend", "Pending recoveries", 0, lambda r: f"={P.a('cumdef', r)}*(1-{I['sev']})-{P.a('cumrec', r)}")
    P.add("val", "Collateral value (end + pending)", lambda r: f"={P.a('t_end', r)}+{P.a('pend', r)}",
          lambda r: f"={P.a('t_end', r)}+{P.a('pend', r)}")
    P.add("cash", "Collateral cash in", None,
          lambda r: f"={P.a('t_int', r)}+{P.a('t_sch', r)}+{P.a('t_pre', r)}+{P.a('rec', r)}")
    P.add("recon", "Balance recon (beg-def-sch-pre-end)", None,
          lambda r: f"={P.a('t_beg', r)}-{P.a('t_def', r)}-{P.a('t_sch', r)}-{P.a('t_pre', r)}-{P.a('t_end', r)}", "0.000000")

    # Tests + Waterfall (mutually referencing, so define column names first)
    W = Series(wb, "Waterfall", "Interest and principal priority of payments, turbo, reserve, write-downs.")
    Tt = Series(wb, "Tests", "OC and IC tests on pro forma balances, breach flags, excess spread.")
    W.add("per", "Period", lambda r: per(r), lambda r: per(r), "0")
    W.add("fee_acc", "Senior fees accrued", None, lambda r: f"={I['fee_rate']}/12*{P.x('t_beg', r)}")
    W.add("fee_due", "Senior fees due", None, lambda r: f"={W.a('fee_acc', r)}+{W.a('fee_short', r - 1)}")
    for k in TR:
        K = k.upper()
        W.add(f"beg_{k}", f"{K} beg balance", None, lambda r, k=k: f"={W.a(f'end_{k}', r - 1)}")
    W.add("cpn_a", "A coupon", None, lambda r: f"=MAX(0,{P.x('fwd', r)}+{I['margin_a']})", PCT)
    W.add("cpn_b", "B coupon", None, lambda r: f"=MAX(0,{P.x('fwd', r)}+{I['margin_b']})", PCT)
    W.add("cpn_c", "C coupon", None, lambda r: f"={I['fixed_c']}", PCT)
    for k in TR:
        K = k.upper()
        W.add(f"acc_{k}", f"{K} interest accrued", None, lambda r, k=k: f"={W.a(f'beg_{k}', r)}*{W.a(f'cpn_{k}', r)}/12")
        W.add(f"due_{k}", f"{K} interest due (incl carry)", None,
              lambda r, k=k: f"={W.a(f'acc_{k}', r)}+{W.a(f'short_{k}', r - 1)}")
    W.add("int_coll", "Interest collections", None, lambda r: f"={P.x('t_int', r)}")
    W.add("res_beg", "Reserve beg", None, lambda r: f"={W.a('res_end', r - 1)}")
    W.add("draw", "Reserve draw (fees, A and B interest left unpaid by collections)", None,
          lambda r: f"=MIN({W.a('res_beg', r)},({W.a('fee_due', r)}-{W.a('fees_c', r)})+({W.a('due_a', r)}-{W.a('paid_a_c', r)})+({W.a('due_b', r)}-{W.a('paid_b_c', r)}))")
    W.add("aif", "Available interest funds", None, lambda r: f"={W.a('int_coll', r)}+{W.a('draw', r)}")
    W.add("apf", "Available principal funds", None, lambda r: f"={P.x('t_sch', r)}+{P.x('t_pre', r)}+{P.x('rec', r)}")
    W.add("seq", "Sequential flag (latches after the first breach)", 0,
          lambda r: f"=IF(OR({I['mode_seq']}=1,{Tt.x('any', r - 1)}=1,{W.a('seq', r - 1)}=1),1,0)", "0")
    W.add("sum_beg", "Sum A+B+C beg", None, lambda r: f"={W.a('beg_a', r)}+{W.a('beg_b', r)}+{W.a('beg_c', r)}")

    def pro(k, r):
        return f"IF({W.a('sum_beg', r)}<=0,0,MIN({W.a('apf', r)},{W.a('sum_beg', r)})*{W.a(f'beg_{k}', r)}/{W.a('sum_beg', r)})"

    W.add("reg_a", "A regular principal", None,
          lambda r: f"=IF({W.a('seq', r)}=1,MIN({W.a('beg_a', r)},{W.a('apf', r)}),{pro('a', r)})")
    W.add("reg_b", "B regular principal", None,
          lambda r: f"=IF({W.a('seq', r)}=1,MIN({W.a('beg_b', r)},{W.a('apf', r)}-{W.a('reg_a', r)}),{pro('b', r)})")
    W.add("reg_c", "C regular principal", None,
          lambda r: f"=IF({W.a('seq', r)}=1,MIN({W.a('beg_c', r)},{W.a('apf', r)}-{W.a('reg_a', r)}-{W.a('reg_b', r)}),{pro('c', r)})")
    W.add("r_prin", "Principal to residual", None,
          lambda r: f"={W.a('apf', r)}-{W.a('reg_a', r)}-{W.a('reg_b', r)}-{W.a('reg_c', r)}")
    for k in TR:
        W.add(f"pf_{k}", f"{k.upper()} pro forma balance", None, lambda r, k=k: f"={W.a(f'beg_{k}', r)}-{W.a(f'reg_{k}', r)}")
    W.add("pf_tot", "Pro forma A+B+C", None, lambda r: f"={W.a('pf_a', r)}+{W.a('pf_b', r)}+{W.a('pf_c', r)}")
    # interest priority of payments
    # steps 1 to 10 run on interest collections only; the reserve draw is allocated afterwards
    W.add("fees_c", "1 Senior fees paid from collections", None, lambda r: f"=MIN({W.a('fee_due', r)},{W.a('int_coll', r)})")
    W.add("rem1", "Remaining after fees", None, lambda r: f"={W.a('int_coll', r)}-{W.a('fees_c', r)}")
    W.add("paid_a_c", "2 A interest paid from collections", None, lambda r: f"=MIN({W.a('due_a', r)},{W.a('rem1', r)})")
    W.add("rem2", "Remaining after A int", None, lambda r: f"={W.a('rem1', r)}-{W.a('paid_a_c', r)}")
    W.add("t1", "3 Turbo (A test breach)", None, lambda r: f"=IF({Tt.x('br_a', r)}=1,MIN({W.a('rem2', r)},{W.a('pf_tot', r)}),0)")
    W.add("rem3", "Remaining after turbo 1", None, lambda r: f"={W.a('rem2', r)}-{W.a('t1', r)}")
    W.add("paid_b_c", "4 B interest paid from collections", None, lambda r: f"=MIN({W.a('due_b', r)},{W.a('rem3', r)})")
    W.add("rem4", "Remaining after B int", None, lambda r: f"={W.a('rem3', r)}-{W.a('paid_b_c', r)}")
    W.add("t2", "5 Turbo (B test breach)", None,
          lambda r: f"=IF({Tt.x('br_b', r)}=1,MIN({W.a('rem4', r)},{W.a('pf_tot', r)}-{W.a('t1', r)}),0)")
    W.add("rem5", "Remaining after turbo 2", None, lambda r: f"={W.a('rem4', r)}-{W.a('t2', r)}")
    W.add("paid_c_c", "6 C interest paid from collections", None, lambda r: f"=MIN({W.a('due_c', r)},{W.a('rem5', r)})")
    W.add("rem6", "Remaining after C int", None, lambda r: f"={W.a('rem5', r)}-{W.a('paid_c_c', r)}")
    W.add("t3", "7 Turbo (C test breach)", None,
          lambda r: f"=IF({Tt.x('br_c', r)}=1,MIN({W.a('rem6', r)},{W.a('pf_tot', r)}-{W.a('t1', r)}-{W.a('t2', r)}),0)")
    W.add("rem7", "Remaining after turbo 3", None, lambda r: f"={W.a('rem6', r)}-{W.a('t3', r)}")
    W.add("def_pre", "Deficiency before turbo 4", None,
          lambda r: f"=MAX(0,{W.a('pf_tot', r)}-{W.a('t1', r)}-{W.a('t2', r)}-{W.a('t3', r)}-{P.x('val', r)})")
    W.add("t4", "8 Turbo (cure deficiency)", None, lambda r: f"=MIN({W.a('rem7', r)},{W.a('def_pre', r)})")
    W.add("rem8", "Remaining after turbo 4", None, lambda r: f"={W.a('rem7', r)}-{W.a('t4', r)}")
    W.add("d_fee", "Reserve draw applied to fees", None,
          lambda r: f"=MIN({W.a('draw', r)},{W.a('fee_due', r)}-{W.a('fees_c', r)})")
    W.add("d_a", "Reserve draw applied to A interest", None,
          lambda r: f"=MIN({W.a('draw', r)}-{W.a('d_fee', r)},{W.a('due_a', r)}-{W.a('paid_a_c', r)})")
    W.add("d_b", "Reserve draw applied to B interest", None,
          lambda r: f"=MIN({W.a('due_b', r)}-{W.a('paid_b_c', r)},{W.a('draw', r)}-{W.a('d_fee', r)}-{W.a('d_a', r)})")
    W.add("turbo", "Turbo total", None,
          lambda r: f"={W.a('t1', r)}+{W.a('t2', r)}+{W.a('t3', r)}+{W.a('t4', r)}")
    W.add("turbo_a", "Turbo to A", None, lambda r: f"=MIN({W.a('pf_a', r)},{W.a('turbo', r)})")
    W.add("turbo_b", "Turbo to B", None, lambda r: f"=MIN({W.a('pf_b', r)},{W.a('turbo', r)}-{W.a('turbo_a', r)})")
    W.add("turbo_c", "Turbo to C", None,
          lambda r: f"=MIN({W.a('pf_c', r)},{W.a('turbo', r)}-{W.a('turbo_a', r)}-{W.a('turbo_b', r)})")
    for k in TR:
        W.add(f"post_{k}", f"{k.upper()} after turbo", None, lambda r, k=k: f"={W.a(f'pf_{k}', r)}-{W.a(f'turbo_{k}', r)}")
    W.add("post_tot", "A+B+C after turbo", None, lambda r: f"={W.a('post_a', r)}+{W.a('post_b', r)}+{W.a('post_c', r)}")
    W.add("defic", "Principal deficiency", None, lambda r: f"=MAX(0,{W.a('post_tot', r)}-{P.x('val', r)})")
    W.add("wd_c", "C write-down", None, lambda r: f"=MIN({W.a('post_c', r)},{W.a('defic', r)})")
    W.add("wd_b", "B write-down", None, lambda r: f"=MIN({W.a('post_b', r)},{W.a('defic', r)}-{W.a('wd_c', r)})")
    W.add("wd_a", "A write-down", None,
          lambda r: f"=MIN({W.a('post_a', r)},{W.a('defic', r)}-{W.a('wd_c', r)}-{W.a('wd_b', r)})")
    for k, bal in (("a", "bal_a"), ("b", "bal_b"), ("c", "bal_c")):
        W.add(f"end_{k}", f"{k.upper()} end balance", f"={I[bal]}",
              lambda r, k=k: f"={W.a(f'post_{k}', r)}-{W.a(f'wd_{k}', r)}")
    W.add("retired", "Notes retired flag", None,
          lambda r: f"=IF({W.a('end_a', r)}+{W.a('end_b', r)}+{W.a('end_c', r)}<0.005,1,0)", "0")
    W.add("res_after", "Reserve after draw", None, lambda r: f"={W.a('res_beg', r)}-{W.a('draw', r)}")
    W.add("topup", "9 Reserve top-up", None,
          lambda r: f"=IF({W.a('retired', r)}=1,0,MIN({W.a('rem8', r)},MAX(0,{I['res_target']}-{W.a('res_after', r)})))")
    W.add("release", "Reserve release", None, lambda r: f"=IF({W.a('retired', r)}=1,{W.a('res_after', r)},0)")
    # once the notes retire, principal collections and the reserve release settle carried fees and A, B, C interest before the residual
    W.add("sweep_avail", "Arrears sweep capacity (principal to residual + release, once retired)", None,
          lambda r: f"=IF({W.a('retired', r)}=1,{W.a('r_prin', r)}+{W.a('release', r)},0)")
    W.add("arr_fee", "Carried fees paid from sweep", None,
          lambda r: f"=MIN({W.a('sweep_avail', r)},{W.a('fee_due', r)}-{W.a('fees_c', r)}-{W.a('d_fee', r)})")
    W.add("arr_a", "Carried A interest paid from sweep", None,
          lambda r: f"=MIN({W.a('sweep_avail', r)}-{W.a('arr_fee', r)},{W.a('due_a', r)}-{W.a('paid_a_c', r)}-{W.a('d_a', r)})")
    W.add("arr_b", "Carried B interest paid from sweep", None,
          lambda r: f"=MIN({W.a('sweep_avail', r)}-{W.a('arr_fee', r)}-{W.a('arr_a', r)},{W.a('due_b', r)}-{W.a('paid_b_c', r)}-{W.a('d_b', r)})")
    W.add("arr_c", "Carried C interest paid from sweep", None,
          lambda r: f"=MIN({W.a('sweep_avail', r)}-{W.a('arr_fee', r)}-{W.a('arr_a', r)}-{W.a('arr_b', r)},{W.a('due_c', r)}-{W.a('paid_c_c', r)})")
    W.add("sweep", "Arrears sweep total", None,
          lambda r: f"={W.a('arr_fee', r)}+{W.a('arr_a', r)}+{W.a('arr_b', r)}+{W.a('arr_c', r)}")
    W.add("r_prin_net", "Principal to residual after arrears", None,
          lambda r: f"={W.a('r_prin', r)}-MIN({W.a('r_prin', r)},{W.a('sweep', r)})")
    W.add("release_r", "Reserve release to residual after arrears", None,
          lambda r: f"={W.a('release', r)}-({W.a('sweep', r)}-MIN({W.a('r_prin', r)},{W.a('sweep', r)}))")
    W.add("fees_paid", "Senior fees paid (collections + reserve + sweep)", None,
          lambda r: f"={W.a('fees_c', r)}+{W.a('d_fee', r)}+{W.a('arr_fee', r)}")
    W.add("paid_a", "A interest paid (collections + reserve + sweep)", None,
          lambda r: f"={W.a('paid_a_c', r)}+{W.a('d_a', r)}+{W.a('arr_a', r)}")
    W.add("paid_b", "B interest paid (collections + reserve + sweep)", None,
          lambda r: f"={W.a('paid_b_c', r)}+{W.a('d_b', r)}+{W.a('arr_b', r)}")
    W.add("paid_c", "C interest paid (collections + sweep)", None,
          lambda r: f"={W.a('paid_c_c', r)}+{W.a('arr_c', r)}")
    W.add("res_end", "Reserve end", f"={I['res_init']}",
          lambda r: f"={W.a('res_after', r)}+{W.a('topup', r)}-{W.a('release', r)}")
    W.add("resid_int", "10 Residual interest distribution", None, lambda r: f"={W.a('rem8', r)}-{W.a('topup', r)}")
    W.add("resid_cash", "Residual cash (int + prin + release)", None,
          lambda r: f"={W.a('resid_int', r)}+{W.a('r_prin_net', r)}+{W.a('release_r', r)}")
    for k in TR:
        W.add(f"prin_{k}", f"{k.upper()} principal paid", None,
              lambda r, k=k: f"={W.a(f'reg_{k}', r)}+{W.a(f'turbo_{k}', r)}")
    W.add("notes_paid", "Notes P+I paid", None,
          lambda r: f"={W.a('paid_a', r)}+{W.a('paid_b', r)}+{W.a('paid_c', r)}+{W.a('prin_a', r)}+{W.a('prin_b', r)}+{W.a('prin_c', r)}")
    for k in TR:
        W.add(f"short_{k}", f"{k.upper()} interest shortfall carried", 0,
              lambda r, k=k: f"={W.a(f'due_{k}', r)}-{W.a(f'paid_{k}', r)}")
    W.add("fee_short", "Fee shortfall carried", 0, lambda r: f"={W.a('fee_due', r)}-{W.a('fees_paid', r)}")
    W.add("cash_chk", "Cash conservation check", None,
          lambda r: f"={W.a('fees_paid', r)}+{W.a('notes_paid', r)}+{W.a('resid_cash', r)}+({W.a('res_end', r)}-{W.a('res_beg', r)})-{P.x('cash', r)}",
          "0.000000")

    # Tests sheet
    Tt.add("per", "Period", lambda r: per(r), lambda r: per(r), "0")
    Tt.add("val", "Collateral value", None, lambda r: f"={P.x('val', r)}")
    Tt.add("num_ic", "IC numerator (interest less fees)", None,
           lambda r: f"=MAX(0,{W.x('int_coll', r)}-{W.x('fee_acc', r)})")
    dens_oc = {"a": lambda r: f"{W.x('pf_a', r)}", "b": lambda r: f"({W.x('pf_a', r)}+{W.x('pf_b', r)})",
               "c": lambda r: f"({W.x('pf_a', r)}+{W.x('pf_b', r)}+{W.x('pf_c', r)})"}
    dens_ic = {"a": lambda r: f"{W.x('acc_a', r)}", "b": lambda r: f"({W.x('acc_a', r)}+{W.x('acc_b', r)})",
               "c": lambda r: f"({W.x('acc_a', r)}+{W.x('acc_b', r)}+{W.x('acc_c', r)})"}
    for k in TR:
        K = k.upper()
        Tt.add(f"oc_{k}", f"OC ratio {K}", None,
               lambda r, k=k: f"=IF({dens_oc[k](r)}<=0.005,999,{Tt.a('val', r)}/{dens_oc[k](r)})", RATIO)
        Tt.add(f"ic_{k}", f"IC ratio {K}", None,
               lambda r, k=k: f"=IF({dens_ic[k](r)}<=0.005,999,{Tt.a('num_ic', r)}/{dens_ic[k](r)})", RATIO)
        Tt.add(f"br_{k}", f"Breach {K} (1=yes)", None,
               lambda r, k=k: f"=IF(OR({Tt.a(f'oc_{k}', r)}<{I[f'oc_{k}']},{Tt.a(f'ic_{k}', r)}<{I[f'ic_{k}']}),1,0)", "0")
    Tt.add("any", "Any breach", 0, lambda r: f"=IF({Tt.a('br_a', r)}+{Tt.a('br_b', r)}+{Tt.a('br_c', r)}>0,1,0)", "0")
    Tt.add("es", "Excess spread ($)", None,
           lambda r: f"={W.x('int_coll', r)}-{W.x('fee_acc', r)}-{W.x('acc_a', r)}-{W.x('acc_b', r)}-{W.x('acc_c', r)}")
    Tt.add("es_bps", "Excess spread (annualized bps of pool)", None,
           lambda r: f"=IF({P.x('t_beg', r)}<=0,0,{Tt.a('es', r)}*12/{P.x('t_beg', r)}*10000)", "0.00")
    Tt.add("cum_loss_pct", "Cum net loss % of pool", 0, lambda r: f"={P.x('cumloss', r)}/{I['pool_bal']}", "0.0000%")
    Tt.add("turbo", "Excess spread diverted (turbo)", None, lambda r: f"={W.x('turbo', r)}")
    Tt.add("res_bal", "Reserve balance", lambda r: f"={W.x('res_end', r)}", lambda r: f"={W.x('res_end', r)}")

    # Cash flows per tranche
    CF = Series(wb, "CashFlows", "Per-tranche cash flows. Period 0 is the purchase at par (negative).")
    CF.add("per", "Period", lambda r: per(r), lambda r: per(r), "0")
    for k in TR:
        K = k.upper()
        CF.add(f"int_{k}", f"{K} interest", None, lambda r, k=k: f"={W.x(f'paid_{k}', r)}")
        CF.add(f"prin_{k}", f"{K} principal", None, lambda r, k=k: f"={W.x(f'prin_{k}', r)}")
        CF.add(f"wd_{k}", f"{K} write-down", None, lambda r, k=k: f"={W.x(f'wd_{k}', r)}")
        CF.add(f"bal_{k}", f"{K} end balance", lambda r, k=k: f"={W.x(f'end_{k}', r)}", lambda r, k=k: f"={W.x(f'end_{k}', r)}")
        CF.add(f"cf_{k}", f"{K} cash flow", lambda r, k=k: f"=-{I[f'bal_{k}']}",
               lambda r, k=k: f"={CF.a(f'int_{k}', r)}+{CF.a(f'prin_{k}', r)}")
    CF.add("int_r", "R interest distribution", None, lambda r: f"={W.x('resid_int', r)}")
    CF.add("prin_r", "R principal + reserve release", None, lambda r: f"={W.x('r_prin_net', r)}+{W.x('release_r', r)}")
    CF.add("cf_r", "R cash flow", lambda r: f"=-{I['bal_r']}", lambda r: f"={CF.a('int_r', r)}+{CF.a('prin_r', r)}")
    CF.add("fwd", "Forward index", lambda r: f"={P.x('fwd', r)}", lambda r: f"={P.x('fwd', r)}", PCT)

    # Retention time series
    RC = Series(wb, "Retention_CF", "Retained-position cash flows and financing sweep by retention option.")
    RC.add("per", "Period", lambda r: per(r), lambda r: per(r), "0")
    RC.add("fwd", "Forward index", lambda r: f"={P.x('fwd', r)}", lambda r: f"={P.x('fwd', r)}", PCT)
    for k in ALL:
        RC.add(f"df_{k}", f"DF {k.upper()} (fwd + FV margin)", 1,
               lambda r, k=k: f"={RC.a(f'df_{k}', r - 1)}/(1+({RC.a('fwd', r)}+{I[f'fvm_{k}']})/12)", "0.000000")
    RC.add("df_fin", "DF financing (fwd + spread)", 1,
           lambda r: f"={RC.a('df_fin', r - 1)}/(1+({RC.a('fwd', r)}+{I['fin_spread']})/12)", "0.000000")

    R = wb.create_sheet("Retention")  # summary sheet, filled below; columns of RC reference its cells
    OPTS = ("rrv", "rrh", "rrl", "euv", "eufl")
    OPT_LABEL = {"rrv": "US RR vertical 246.4(a)(1)", "rrh": "US RR horizontal 246.4(a)(2)",
                 "rrl": "US RR L-shaped 246.4(a)(3)", "euv": "EU Art 6(3)(a) vertical",
                 "eufl": "EU Art 6(3)(d) first loss"}
    CLS = ("a", "b", "c", "r1", "r2")  # R1 retained first-loss residual subclass, R2 sold residual subclass
    # layout of Retention sheet
    fv_row = {k: 6 + i for i, k in enumerate(ALL)}  # FV table rows 6..9, column C = FV
    frac_row0 = 15
    frac_row = {o: frac_row0 + i for i, o in enumerate(OPTS)}  # columns B..F = A,B,C,R1,R2; G = R2 par
    frac_col = {k: get_column_letter(2 + i) for i, k in enumerate(CLS)}
    P2_COL, PAR_COL, FV_COL, ADV_COL, PARPCT_COL, FVPCT_COL, REG_COL = "G", "H", "I", "J", "K", "L", "M"
    hz_row = frac_row0 + len(OPTS) + 3  # helper block: whole-class flags for the L-shaped horizontal component
    fr = lambda o, k: f"Retention!${frac_col[k]}${frac_row[o]}"  # noqa: E731
    p2c = lambda o: f"Retention!${P2_COL}${frac_row[o]}"  # noqa: E731
    fvc = lambda k: f"Retention!$C${fv_row[k]}"  # noqa: E731
    parc = lambda k: I[f"bal_{k}"]  # noqa: E731
    cost_cell = lambda o: f"Retention!${PAR_COL}${frac_row[o]}"  # noqa: E731
    adv_cell = lambda o: f"Retention!${ADV_COL}${frac_row[o]}"  # noqa: E731

    RC.add("cf_r_raw", "R residual cash (interest distribution + principal + release)", None,
           lambda r: f"={CF.x('cf_r', r)}")
    for o in OPTS:
        # R2 is a priority claim on residual cash accruing at forward + residual FV margin; R1 gets the rest
        RC.add(f"r2b_{o}", f"{o} R2 claim balance", lambda r, o=o: f"={p2c(o)}",
               lambda r, o=o: f"={RC.a(f'r2b_{o}', r - 1)}*(1+({RC.a('fwd', r)}+{I['fvm_r']})/12)-{RC.a(f'r2p_{o}', r)}")
        RC.add(f"r2p_{o}", f"{o} R2 payment", None,
               lambda r, o=o: f"=MIN({RC.a('cf_r_raw', r)},{RC.a(f'r2b_{o}', r - 1)}*(1+({RC.a('fwd', r)}+{I['fvm_r']})/12))")
        RC.add(f"cf_{o}", f"{OPT_LABEL[o]} retained cash flow", lambda r, o=o: f"=-{cost_cell(o)}",
               lambda r, o=o: "=" + "+".join(f"{fr(o, k)}*{CF.x(f'cf_{k}', r)}" for k in TR)
               + f"+{fr(o, 'r1')}*({RC.a('cf_r_raw', r)}-{RC.a(f'r2p_{o}', r)})+{fr(o, 'r2')}*{RC.a(f'r2p_{o}', r)}")
        RC.add(f"lb_{o}", f"{o} loan balance", lambda r, o=o: f"={adv_cell(o)}",
               lambda r, o=o: f"={RC.a(f'lb_{o}', r - 1)}+{RC.a(f'li_{o}', r)}-{RC.a(f'lp_{o}', r)}")
        RC.add(f"li_{o}", f"{o} loan interest", None,
               lambda r, o=o: f"={RC.a(f'lb_{o}', r - 1)}*({RC.a('fwd', r)}+{I['fin_spread']})/12")
        RC.add(f"lp_{o}", f"{o} loan payment (sweep)", None,
               lambda r, o=o: f"=MIN({RC.a(f'cf_{o}', r)},{RC.a(f'lb_{o}', r - 1)}+{RC.a(f'li_{o}', r)})")
        RC.add(f"eq_{o}", f"{o} sponsor cash flow after loan", lambda r, o=o: f"=-({cost_cell(o)}-{adv_cell(o)})",
               lambda r, o=o: f"={RC.a(f'cf_{o}', r)}-{RC.a(f'lp_{o}', r)}")

    R.cell(1, 1, "Risk retention sizing (US Regulation RR, 17 CFR Part 246; EU Regulation 2017/2402 Article 6)").font = BOLD
    R.cell(2, 1, "Fair values are PV of the selected scenario's tranche cash flows at forward + assumed margin (illustrative). Retained cost basis for yield is par.")
    R.cell(3, 1, "Horizontal options retain R1, a separately identified first-loss residual subclass paid after the sold subclass R2 (17 CFR 246.2, eligible horizontal residual interest). R1 is the whole residual in the other options.")
    R.cell(4, 1, "Fair value of each ABS interest").font = BOLD
    for j, h in enumerate(("Tranche", "Par", "Fair value used for sizing", "Live PV fair value (selected scenario)", "FV used % of par")):
        R.cell(5, 1 + j, h).font = BOLD
    for k in ALL:
        rr = fv_row[k]
        R.cell(rr, 1, k.upper())
        R.cell(rr, 2, f"={parc(k)}").number_format = NUM
        R.cell(rr, 3, f"=IF(ISNUMBER({I[f'fv_lock_{k}']}),{I[f'fv_lock_{k}']},D{rr})").number_format = NUM
        R.cell(rr, 4, f"=SUMPRODUCT({CF.rng(f'cf_{k}', ROW0 + 1)},{RC.rng(f'df_{k}', ROW0 + 1)})").number_format = NUM
        R.cell(rr, 5, f"=C{rr}/B{rr}").number_format = "0.0000%"
    R.cell(10, 1, "Total")
    R.cell(10, 2, "=SUM(B6:B9)").number_format = NUM
    R.cell(10, 3, "=SUM(C6:C9)").number_format = NUM
    R.cell(10, 4, "=SUM(D6:D9)").number_format = NUM
    R.cell(12, 1, "US minimum horizontal fair value (5% of FV of all ABS interests), also the 246.4(b) cash reserve alternative")
    R.cell(12, 3, f"={I['ret_pct']}*C10").number_format = NUM
    R.cell(13, 1, "EU 6(3)(d) required nominal (5% of pool at origination); same amount for 6(3)(c) and 6(3)(e)")
    R.cell(13, 3, f"={I['ret_pct']}*{I['pool_bal']}").number_format = NUM

    for j, h in enumerate(("Option (fraction of each class retained)", "A", "B", "C",
                           "R1 (retained first loss; whole R unless R2 is carved out)", "R2 (sold residual subclass)",
                           "R2 par carved out of R", "Retained par (cost)", "Retained fair value", "Advance",
                           "Par % of pool", "FV % of all FV",
                           "Regulatory measure (246.4(a): vertical %, horizontal FV %, or their sum)")):
        c = R.cell(frac_row0 - 1, 1 + j, h)
        c.font = BOLD
        c.alignment = Alignment(wrap_text=True)

    def whole_flags(req, row):
        """Whole junior notes C, B, A added while the residual alone falls short of req; written to row."""
        rem = f"MAX(0,{req}-$C$9)"
        for k in ("c", "b", "a"):
            col = frac_col[k]
            R[f"{col}{row}"] = f"=IF({rem}>0,1,0)"
            R[f"{col}{row}"].number_format = RATIO
            rem = f"({rem}-{col}{row}*{fvc(k)})"

    for o in OPTS:
        R.cell(frac_row[o], 1, OPT_LABEL[o])
    for k in TR:
        R[f"{frac_col[k]}{frac_row['rrv']}"] = f"={I['ret_pct']}"
        R[f"{frac_col[k]}{frac_row['euv']}"] = f"={I['ret_pct']}*{I[f'eu_sold_{k}']}"
    R[f"{frac_col['r1']}{frac_row['rrv']}"] = f"={I['ret_pct']}"
    R[f"{frac_col['r1']}{frac_row['euv']}"] = f"={I['ret_pct']}*{I['eu_sold_r']}"
    req_h = f"{I['ret_pct']}*$C$10"
    req_l = f"({I['ret_pct']}-{I['lshape_v']})*$C$10/(1-{I['lshape_v']})"  # vertical slice already holds v of R1
    whole_flags(req_h, frac_row["rrh"])
    R[f"{frac_col['r1']}{frac_row['rrh']}"] = f"=IF({req_h}>0,1,0)"
    R[f"{P2_COL}{frac_row['rrh']}"] = f"=MAX(0,$C$9-{req_h})"
    R.cell(hz_row - 1, 1, "Helper: whole-class flags for the horizontal component of the L-shaped option").font = BOLD
    R.cell(hz_row, 1, "L-shape horizontal component whole-class flags")
    whole_flags(req_l, hz_row)
    for k in TR:
        R[f"{frac_col[k]}{frac_row['rrl']}"] = f"=MAX({I['lshape_v']},{frac_col[k]}{hz_row})"
    R[f"{frac_col['r1']}{frac_row['rrl']}"] = f"=IF({req_l}>0,1,0)"
    R[f"{frac_col['r2']}{frac_row['rrl']}"] = f"={I['lshape_v']}"
    R[f"{P2_COL}{frac_row['rrl']}"] = f"=MAX(0,$C$9-{req_l})"
    # EU first loss: whole classes from the bottom until nominal covers 5% of pool
    need = f"{I['ret_pct']}*{I['pool_bal']}"
    R[f"{frac_col['r1']}{frac_row['eufl']}"] = f"=IF(0<{need},1,0)"
    cum = f"{frac_col['r1']}{frac_row['eufl']}*({parc('r')}-{I['res_init']})"  # principal only, reserve cannot absorb losses
    for k in ("c", "b", "a"):
        R[f"{frac_col[k]}{frac_row['eufl']}"] = f"=IF({cum}<{need},1,0)"
        cum = f"{cum}+{frac_col[k]}{frac_row['eufl']}*{parc(k)}"
    for o in OPTS:
        rw = frac_row[o]
        for k in CLS:
            if R[f"{frac_col[k]}{rw}"].value is None:
                R[f"{frac_col[k]}{rw}"] = 0
            R[f"{frac_col[k]}{rw}"].number_format = RATIO
        if R[f"{P2_COL}{rw}"].value is None:
            R[f"{P2_COL}{rw}"] = 0
        R[f"{P2_COL}{rw}"].number_format = NUM
        p2 = f"{P2_COL}{rw}"
        f_ = lambda k: f"{frac_col[k]}{rw}"  # noqa: E731
        notes_par = "+".join(f"{f_(k)}*{parc(k)}" for k in TR)
        notes_fv = "+".join(f"{f_(k)}*$C${fv_row[k]}" for k in TR)
        notes_adv = "+".join(f"{f_(k)}*{parc(k)}*{I[f'adv_{k}']}" for k in TR)
        r1_par = f"({parc('r')}-{p2})"
        R[f"{PAR_COL}{rw}"] = f"={notes_par}+{f_('r1')}*{r1_par}+{f_('r2')}*{p2}"
        R[f"{FV_COL}{rw}"] = f"={notes_fv}+{f_('r1')}*($C$9-{p2})+{f_('r2')}*{p2}"
        R[f"{ADV_COL}{rw}"] = f"={notes_adv}+({f_('r1')}*{r1_par}+{f_('r2')}*{p2})*{I['adv_r']}"
        R[f"{PARPCT_COL}{rw}"] = f"={PAR_COL}{rw}/{I['pool_bal']}"
        R[f"{FVPCT_COL}{rw}"] = f"={FV_COL}{rw}/$C$10"
        for c_ in (PAR_COL, FV_COL, ADV_COL):
            R[f"{c_}{rw}"].number_format = NUM
        R[f"{PARPCT_COL}{rw}"].number_format = "0.0000%"
        R[f"{FVPCT_COL}{rw}"].number_format = "0.0000%"
    R[f"{REG_COL}{frac_row['rrv']}"] = f"={I['ret_pct']}"
    r1_fv = lambda o: f"{frac_col['r1']}{frac_row[o]}*($C$9-{P2_COL}{frac_row[o]})"  # noqa: E731
    hz_fv = lambda o, row: "+".join([r1_fv(o)] + [f"{frac_col[k]}{row}*{fvc(k)}" for k in TR])  # noqa: E731
    R[f"{REG_COL}{frac_row['rrh']}"] = f"=({hz_fv('rrh', frac_row['rrh'])})/$C$10"
    R[f"{REG_COL}{frac_row['rrl']}"] = (
        f"={I['lshape_v']}+(1-{I['lshape_v']})*({hz_fv('rrl', hz_row)})/$C$10")
    for o in ("euv", "eufl"):
        R[f"{REG_COL}{frac_row[o]}"] = "n/a"
    for o in OPTS:
        R[f"{REG_COL}{frac_row[o]}"].number_format = "0.0000%"

    out_row0 = hz_row + 4
    R.cell(out_row0 - 1, 1, "Retained position and financing view (selected scenario)").font = BOLD
    heads = ("Option", "Total cash received", "Yield at cost (monthly IRR x 12)", "Advance", "Advance / cost",
             "PV coverage (retained CF PV at financing rate / advance)", "Loan repaid in month", "Unpaid loan at horizon", "Levered IRR")
    for j, h in enumerate(heads):
        c = R.cell(out_row0, 1 + j, h)
        c.font = BOLD
        c.alignment = Alignment(wrap_text=True)
    out_row = {o: out_row0 + 1 + i for i, o in enumerate(OPTS)}
    last = ROW0 + HORIZON
    for o in OPTS:
        rw = out_row[o]
        R.cell(rw, 1, OPT_LABEL[o])
        R.cell(rw, 2, f"=SUM({RC.rng(f'cf_{o}', ROW0 + 1)})").number_format = NUM
        R.cell(rw, 3, irr12(RC.rng(f"cf_{o}"))).number_format = "0.0000%"
        R.cell(rw, 4, f"={adv_cell(o)}").number_format = NUM
        R.cell(rw, 5, f"=D{rw}/{cost_cell(o)}").number_format = "0.0000%"
        R.cell(rw, 6, f"=SUMPRODUCT({RC.rng(f'cf_{o}', ROW0 + 1)},{RC.rng('df_fin', ROW0 + 1)})/D{rw}").number_format = RATIO
        R.cell(rw, 7, f'=IF({RC.x(f"lb_{o}", last)}<=0.005,COUNTIF({RC.rng(f"lb_{o}", ROW0 + 1)},">0.005")+1,"not repaid")')
        R.cell(rw, 8, f"={RC.x(f'lb_{o}', last)}").number_format = NUM
        R.cell(rw, 9, irr12(RC.rng(f"eq_{o}"))).number_format = "0.0000%"
    R.column_dimensions["A"].width = 60
    for c_ in "BCDEFGHIJKLM":
        R.column_dimensions[c_].width = 20

    # Scenario summary
    S = wb.create_sheet("Scenario_Summary")
    S.cell(1, 1, "Scenario summary and integrity checks (live for the scenario selected on Inputs)").font = BOLD
    S.cell(2, 1, "Active scenario")
    S.cell(2, 2, f"={I['scen_name']}")
    for j, (lab, key) in enumerate((("CPR", "cpr"), ("CDR", "cdr"), ("Severity", "sev"), ("Lag (months)", "lag"), ("Forward shift bps", "shift"))):
        S.cell(3, 1 + j, lab).font = BOLD
        S.cell(4, 1 + j, f"={I[key]}")
    heads = ("Tranche", "Original balance", "Interest paid", "Principal paid", "Write-down / capital shortfall", "Ending balance",
             "WAL (years)", "Yield at par (monthly IRR x 12)", "Principal window first", "Principal window last",
             "Conservation check (orig - paid - writedown - end)", "Write-down / capital loss above USD 1 (1/0)",
             "Unpaid interest at horizon", "Peak unpaid interest (any month)",
             "Any shortfall: write-down or unpaid interest above USD 1 (1/0)")
    for j, h in enumerate(heads):
        c = S.cell(6, 1 + j, h)
        c.font = BOLD
        c.alignment = Alignment(wrap_text=True)
    srow = {k: 7 + i for i, k in enumerate(ALL)}
    pr = lambda k: CF.rng(f"prin_{k}", ROW0 + 1)  # noqa: E731
    for k in ALL:
        rw = srow[k]
        S.cell(rw, 1, k.upper())
        S.cell(rw, 2, f"={I[f'bal_{k}']}").number_format = NUM
        S.cell(rw, 3, f"=SUM({CF.rng(f'int_{k}', ROW0 + 1)})").number_format = NUM
        S.cell(rw, 4, f"=SUM({pr(k)})").number_format = NUM
        if k == "r":
            S.cell(rw, 5, f"=MAX(0,B{rw}-C{rw}-D{rw})").number_format = NUM
            S.cell(rw, 6, 0).number_format = NUM
        else:
            S.cell(rw, 5, f"=SUM({CF.rng(f'wd_{k}', ROW0 + 1)})").number_format = NUM
            S.cell(rw, 6, f"={CF.x(f'bal_{k}', ROW0 + HORIZON)}").number_format = NUM
        S.cell(rw, 7, f'=IF(D{rw}<=0.005,"n/a",SUMPRODUCT({CF.rng("per", ROW0 + 1)},{pr(k)})/12/D{rw})').number_format = "0.0000"
        S.cell(rw, 8, irr12(CF.rng(f"cf_{k}"))).number_format = "0.0000%"
        S.cell(rw, 9, f'=IF(D{rw}<=0.005,"n/a",_xlfn.MINIFS({CF.rng("per", ROW0 + 1)},{pr(k)},">0.005"))')
        S.cell(rw, 10, f'=IF(D{rw}<=0.005,"n/a",_xlfn.MAXIFS({CF.rng("per", ROW0 + 1)},{pr(k)},">0.005"))')
        if k == "r":
            S.cell(rw, 11, "n/a")
        else:
            S.cell(rw, 11, f"=B{rw}-D{rw}-E{rw}-F{rw}").number_format = NUM
        S.cell(rw, 12, f"=IF(E{rw}>1,1,0)")
        if k == "r":
            S.cell(rw, 13, 0).number_format = NUM
            S.cell(rw, 14, 0).number_format = NUM
        else:
            S.cell(rw, 13, f"={W.x(f'short_{k}', ROW0 + HORIZON)}").number_format = NUM
            S.cell(rw, 14, f"=MAX({W.rng(f'short_{k}', ROW0 + 1)})").number_format = NUM
        S.cell(rw, 15, f"=IF(OR(L{rw}=1,N{rw}>1),1,0)")
    S.cell(12, 1, "For R: column C is residual interest distributions, D is principal distributions plus reserve release, E is the capital shortfall (par less cash received, floored at 0).")
    S.cell(14, 1, "Integrity checks (sum of absolute differences, should be 0)").font = BOLD
    checks = (("Pool balance reconciliation", f"=SUMPRODUCT(ABS({P.rng('recon', ROW0 + 1)}))"),
              ("Cash conservation", f"=SUMPRODUCT(ABS({W.rng('cash_chk', ROW0 + 1)}))"),
              ("Notes P+I plus fees paid above collateral cash plus reserve draw and release (sum of excesses)",
               f"=SUMPRODUCT((({W.rng('notes_paid', ROW0 + 1)}+{W.rng('fees_paid', ROW0 + 1)}-{P.rng('cash', ROW0 + 1)}-{W.rng('draw', ROW0 + 1)}-{W.rng('release', ROW0 + 1)})>0.005)*({W.rng('notes_paid', ROW0 + 1)}+{W.rng('fees_paid', ROW0 + 1)}-{P.rng('cash', ROW0 + 1)}-{W.rng('draw', ROW0 + 1)}-{W.rng('release', ROW0 + 1)}))"),
              ("Tranche conservation A", "=ABS(K7)"), ("Tranche conservation B", "=ABS(K8)"), ("Tranche conservation C", "=ABS(K9)"),
              ("Months with a trigger breach", f"=SUM({Tt.rng('any', ROW0 + 1)})"),
              ("Excess spread at month 12 (bps annualized)", f"={Tt.x('es_bps', ROW0 + 12)}"),
              ("Cumulative net loss % of pool at horizon", f"={Tt.x('cum_loss_pct', ROW0 + HORIZON)}"),
              ("Reserve draw not applied to fees, A or B interest (sum of absolute differences, should be 0)",
               f"=SUMPRODUCT(ABS({W.rng('draw', ROW0 + 1)}-{W.rng('d_fee', ROW0 + 1)}-{W.rng('d_a', ROW0 + 1)}-{W.rng('d_b', ROW0 + 1)}))"))
    for i, (lab, f) in enumerate(checks):
        S.cell(15 + i, 1, lab)
        S.cell(15 + i, 2, f).number_format = "0.000000"
    S.cell(29, 1, "Class labels: write-down only versus write-down or unpaid interest").font = BOLD
    def _pick(order, col):
        expr = '"none"'
        for k in reversed(order):
            expr = f'IF({col}{srow[k]}=1,"{k.upper()}",{expr})'
        return "=" + expr
    label_rows = (("Most senior class with a principal write-down (interest deferrals ignored)", _pick(ALL, "L")),
                  ("Most junior class with a principal write-down (first to absorb a loss)", _pick(tuple(reversed(ALL)), "L")),
                  ("Most senior class with any shortfall (write-down or unpaid interest)", _pick(ALL, "O")),
                  ("First class with any shortfall (most junior class with a write-down or unpaid interest)",
                   _pick(tuple(reversed(ALL)), "O")))
    for i, (lab, fml) in enumerate(label_rows):
        S.cell(30 + i, 1, lab)
        S.cell(30 + i, 2, fml)
    S.cell(26, 1, "Break-even CDR per tranche is solved by bisection in python/ (Excel has no live solver). Set the CDR of a scenario to a Python break-even value here to confirm zero loss just below it.")
    S.cell(27, 1, "Scenario table lives on Inputs; change the selected scenario number there to recalculate every sheet.")
    S.column_dimensions["A"].width = 60
    for j in range(2, 16):
        S.column_dimensions[get_column_letter(j)].width = 18

    for s in (P, W, Tt, CF, RC):
        s.fill()

    # Defined names for reconciliation
    def name(nm, ref):
        wb.defined_names[nm] = DefinedName(nm, attr_text=ref)

    name("scen_sel", I["scen_sel"])
    name("SCEN_TABLE", f"Inputs!$A${s1}:$F${s2}")
    for k in ALL:
        K = k.upper()
        name(f"WAL_{K}", f"Scenario_Summary!$G${srow[k]}")
        name(f"YIELD_{K}", f"Scenario_Summary!$H${srow[k]}")
        name(f"LOSS_{K}", f"Scenario_Summary!$E${srow[k]}")
        name(f"PRINPAID_{K}", f"Scenario_Summary!$D${srow[k]}")
        name(f"INTPAID_{K}", f"Scenario_Summary!$C${srow[k]}")
        name(f"CF_{K}", CF.rng(f"cf_{k}"))
        name(f"FV_{K}", f"Retention!$C${fv_row[k]}")
        name(f"FVLIVE_{K}", f"Retention!$D${fv_row[k]}")
        name(f"FVLOCK_{K}", I[f"fv_lock_{k}"])
    name("ENDBAL_A", CF.rng("bal_a"))
    name("ENDBAL_B", CF.rng("bal_b"))
    name("ENDBAL_C", CF.rng("bal_c"))
    name("CHECK_POOL_RECON", "Scenario_Summary!$B$15")
    name("CHECK_CASH", "Scenario_Summary!$B$16")
    name("CHECK_NOTES_OVER_CASH", "Scenario_Summary!$B$17")
    name("BREACH_MONTHS", "Scenario_Summary!$B$21")
    name("CHECK_RESERVE_ALLOC", "Scenario_Summary!$B$24")
    name("LABELS", "Scenario_Summary!$B$30:$B$33")
    for k in ALL:
        K = k.upper()
        name(f"UNPAID_END_{K}", f"Scenario_Summary!$M${srow[k]}")
        name(f"UNPAID_PEAK_{K}", f"Scenario_Summary!$N${srow[k]}")
        name(f"ANYSHORT_{K}", f"Scenario_Summary!$O${srow[k]}")
    name("FRAC_TABLE", f"Retention!$B${frac_row['rrv']}:${P2_COL}${frac_row['eufl']}")
    name("RET_AMOUNTS", f"Retention!${PAR_COL}${frac_row['rrv']}:${ADV_COL}${frac_row['eufl']}")
    name("RET_REG", f"Retention!${REG_COL}${frac_row['rrv']}:${REG_COL}${frac_row['eufl']}")
    name("RET_OUT", f"Retention!$B${out_row['rrv']}:$I${out_row['eufl']}")
    name("INPUT_POOL_BAL", I["pool_bal"])
    name("INPUT_BAL_R", I["bal_r"])
    for o in OPTS:
        name(f"RETCF_{o.upper()}", RC.rng(f"cf_{o}"))
    name("POOL_END", P.rng("t_end"))
    name("ARREARS_SWEEP", W.rng("sweep"))
    name("RESID_CASH", W.rng("resid_cash"))
    name("POOL_CASH", P.rng("cash"))
    name("ES_BPS", Tt.rng("es_bps"))
    for kk in ("scen_sel", "pool_bal", "bal_r", "mode", "oc_a", "oc_b", "oc_c", "ic_a", "ic_b", "ic_c", "res_init",
               "res_target", "margin_a", "margin_b", "fixed_c", "svc", "trustee", "fwd_start", "fwd_end", "fwd_tau",
               "bal_a", "bal_b", "bal_c", "ret_pct", "lshape_v", "fin_spread"):
        name(f"IN_{kk.upper()}", I[kk])

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def embed_cached_values(path):
    """Round-trip through LibreOffice so the file carries calculated values for viewers that do not recalculate."""
    exe = next((c for c in SOFFICE_CANDIDATES if Path(c).exists() or shutil.which(c)), None)
    if exe is None:
        return False
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run([exe, "-headless", "-convert-to", "xlsx", "-outdir", tmp, str(path)], check=True,
                       capture_output=True, timeout=180)
        shutil.copyfile(Path(tmp) / Path(path).name, path)
    return True


if __name__ == "__main__":
    out = build()
    print("wrote", out, "| cached values embedded:", embed_cached_values(out))
