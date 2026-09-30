"""Recalculate the Excel workbook with LibreOffice and reconcile it to the Python model.

Run: py -3 -m waterfall.reconcile_excel   (from the python/ folder, after building the workbook)
"""
import json
import math
import shutil
import subprocess
import tempfile
import time
from dataclasses import replace
from pathlib import Path

import numpy as np
import openpyxl

from . import columns as C
from .config import (BASE_SCENARIO, EXCEL_PATH, NOTES, OUTPUT_DIR, SCENARIOS, SOFFICE_CANDIDATES, TRANCHES, Scenario, Deal)
from .metrics import tranche_cash_flows, tranche_summary
from .retention import OPTIONS, RET_CLASSES, fair_values, retention_amounts, retention_report, retention_structure
from .scenarios import breakeven_cdr, class_labels
from .waterfall import run_waterfall

# Tolerances, stated in the README
TOL_USD = 0.01  # absolute, per cash flow cell and per dollar total
TOL_RATE = 1e-6  # absolute, for WAL (years) and annualized yields
TOL_LEVERED = 1e-5  # levered IRR is more sensitive to the IRR solver
TOL_RATIO = 1e-9  # for fractions retained
EXTRA_SCENARIOS = (Scenario("Late-default repro", 0.02, 0.18, 0.20, 12), Scenario("Severity 70 CDR 40", 0.10, 0.40, 0.70, 6))


def find_soffice():
    for cand in SOFFICE_CANDIDATES:
        if Path(cand).exists() or shutil.which(cand):
            return cand
    return None


def _split_ref(wb, defined_name):
    ref = wb.defined_names[defined_name].attr_text
    sheet, rng = ref.split("!")
    return sheet.strip("'"), rng.replace("$", "")


def _flat(cells):
    if not isinstance(cells, tuple):
        return [cells.value]
    out = []
    for row in cells:
        out.extend(_flat(row) if isinstance(row, tuple) else [row.value])
    return out


class Book:
    """Values from a recalculated workbook, addressed through the original workbook's defined names."""

    def __init__(self, orig: Path, recalculated: Path):
        self.names = openpyxl.load_workbook(orig)
        self.vals = openpyxl.load_workbook(recalculated, data_only=True)

    def get(self, name):
        sheet, rng = _split_ref(self.names, name)
        return _flat(self.vals[sheet][rng])

    def num(self, name):
        return _num(self.get(name)[0])


def recalc(src: Path, outdir: Path) -> Path:
    exe = find_soffice()
    outdir.mkdir(parents=True, exist_ok=True)
    for attempt in range(3):  # LibreOffice occasionally fails right after a previous headless run
        proc = subprocess.run([exe, "-headless", "-convert-to", "xlsx", "-outdir", str(outdir), str(src)],
                              capture_output=True, timeout=180)
        if proc.returncode == 0 and (outdir / src.name).exists():
            return outdir / src.name
        time.sleep(2)
    raise RuntimeError(f"LibreOffice recalculation failed: {proc.stderr.decode(errors='replace')}")


def prepare(src: Path, dst: Path, scen_index: int, locks: dict = None, scen_override: tuple = None,
            mode: str = None):
    wb = openpyxl.load_workbook(src)
    sheet, addr = _split_ref(wb, "scen_sel")
    wb[sheet][addr] = scen_index
    if mode:
        sheet, addr = _split_ref(wb, "IN_MODE")
        wb[sheet][addr] = mode
    for k, v in (locks or {}).items():
        sheet, addr = _split_ref(wb, f"FVLOCK_{k}")
        wb[sheet][addr] = v
    if scen_override:
        sheet, rng = _split_ref(wb, "SCEN_TABLE")
        row = wb[sheet][rng][scen_index - 1]
        for cell, v in zip(row, scen_override):
            cell.value = v
    wb.save(dst)


def _num(v):
    return float("nan") if not isinstance(v, (int, float)) else float(v)


def _diff(a, b):
    """Absolute difference; two nans count as equal, one nan as infinite."""
    a_nan, b_nan = math.isnan(a), math.isnan(b)
    if a_nan and b_nan:
        return 0.0
    if a_nan or b_nan:
        return float("inf")
    return abs(a - b)


def compare(book: Book, deal: Deal, scen, base_fv: dict, locked: bool) -> dict:
    res = run_waterfall(deal, scen)
    summ = tranche_summary(res)
    worst = {"cash_flow_usd": 0.0, "usd_totals": 0.0, "wal_years": 0.0, "yield": 0.0, "fair_value_usd": 0.0,
             "retention_usd": 0.0, "retention_fractions": 0.0, "retention_rates": 0.0, "levered_irr": 0.0,
             "unpaid_interest_usd": 0.0, "arrears_sweep_usd": 0.0, "residual_cash_usd": 0.0}
    flags_ok = True
    detail = {}
    for x in TRANCHES:
        cf = tranche_cash_flows(res, x)["cf"]
        xl = np.array([_num(v) for v in book.get(f"CF_{x}")])
        worst["cash_flow_usd"] = max(worst["cash_flow_usd"], float(np.abs(xl - cf).max()))
        s = summ[x]
        for nm, key in ((f"INTPAID_{x}", "interest_paid"), (f"PRINPAID_{x}", "principal_paid"), (f"LOSS_{x}", "writedown")):
            worst["usd_totals"] = max(worst["usd_totals"], _diff(book.num(nm), s[key]))
        worst["wal_years"] = max(worst["wal_years"], _diff(book.num(f"WAL_{x}"), s["wal_years"]))
        worst["yield"] = max(worst["yield"], _diff(book.num(f"YIELD_{x}"), s["yield_at_par"]))
        worst["unpaid_interest_usd"] = max(worst["unpaid_interest_usd"],
                                           _diff(book.num(f"UNPAID_END_{x}"), s["unpaid_interest_end"]),
                                           _diff(book.num(f"UNPAID_PEAK_{x}"), s["unpaid_interest_peak"]))
        flags_ok = flags_ok and int(book.num(f"ANYSHORT_{x}")) == int(s["any_shortfall"])
        detail[x] = {"excel_wal": book.num(f"WAL_{x}"), "py_wal": s["wal_years"],
                     "excel_yield": book.num(f"YIELD_{x}"), "py_yield": s["yield_at_par"],
                     "excel_loss": book.num(f"LOSS_{x}"), "py_loss": s["writedown"]}
    sweep_py = res[C.ARR_FEE] + sum(res[C.tr(x, C.ARR)] for x in NOTES)
    worst["arrears_sweep_usd"] = float(np.abs(np.array([_num(v) for v in book.get("ARREARS_SWEEP")][1:]) - sweep_py[1:]).max())
    worst["residual_cash_usd"] = float(np.abs(np.array([_num(v) for v in book.get("RESID_CASH")][1:]) - res[C.RESID_CASH][1:]).max())
    live_fv = fair_values(res)
    for x in TRANCHES:
        worst["fair_value_usd"] = max(worst["fair_value_usd"], _diff(book.num(f"FVLIVE_{x}"), live_fv[x]))
    # Python sizes on base fair values (locked in Excel) or on this scenario (unlocked base run)
    sizing_fv = base_fv if locked else live_fv
    am = retention_amounts(deal, sizing_fv)
    structure = retention_structure(deal, sizing_fv)
    n_opt, n_cls = len(OPTIONS), len(RET_CLASSES)
    frac_xl = np.array([_num(v) for v in book.get("FRAC_TABLE")]).reshape(n_opt, n_cls + 1)  # last column is R2 par
    reg_xl = [_num(v) for v in book.get("RET_REG")]
    amt_xl = np.array([_num(v) for v in book.get("RET_AMOUNTS")]).reshape(n_opt, 3)
    out_xl = np.array(book.get("RET_OUT"), dtype=object).reshape(n_opt, 8)
    rep = retention_report(res, structure)
    for i, o in enumerate(OPTIONS):
        a, f = am["options"][o], rep[o]["financing"]
        py_frac = np.array([a["fractions"][x] for x in RET_CLASSES])
        worst["retention_fractions"] = max(worst["retention_fractions"],
                                           float(np.abs(frac_xl[i][:n_cls] - py_frac).max()))
        if "regulatory_measure" in a:
            worst["retention_rates"] = max(worst["retention_rates"], _diff(reg_xl[i], a["regulatory_measure"]))
        worst["retention_usd"] = max(
            worst["retention_usd"], abs(frac_xl[i][n_cls] - a["r2_par"]),
            abs(amt_xl[i][0] - a["retained_par"]), abs(amt_xl[i][1] - a["retained_fair_value"]),
            abs(amt_xl[i][2] - f["advance"]), abs(_num(out_xl[i][0]) - rep[o]["total_cash"]),
            abs(_num(out_xl[i][6]) - f["unpaid_at_horizon"]))
        worst["retention_rates"] = max(worst["retention_rates"], _diff(_num(out_xl[i][1]), rep[o]["yield_at_cost"]),
                                       _diff(_num(out_xl[i][4]), f["pv_coverage"]))
        worst["levered_irr"] = max(worst["levered_irr"], _diff(_num(out_xl[i][7]), f["levered_irr"]))
        xl_month, py_month = out_xl[i][5], f["repaid_month"]
        if (xl_month == "not repaid") != (py_month is None) or (py_month is not None and int(xl_month) != py_month):
            worst["retention_rates"] = float("inf")
    labels_xl = book.get("LABELS")
    labels_ok = labels_xl == list(class_labels(summ).values())
    checks = {"reserve_draw_alloc": book.num("CHECK_RESERVE_ALLOC"), "labels_match": labels_ok, "any_shortfall_flags_match": flags_ok,
              "pool_recon_sum_abs": book.num("CHECK_POOL_RECON"), "cash_conservation_sum_abs": book.num("CHECK_CASH"),
              "notes_over_cash_sum": book.num("CHECK_NOTES_OVER_CASH"), "breach_months_excel": book.num("BREACH_MONTHS"),
              "breach_months_python": float(res[C.BREACH_ANY].sum())}
    ok = (worst["cash_flow_usd"] < TOL_USD and worst["usd_totals"] < TOL_USD and worst["wal_years"] < TOL_RATE
          and worst["yield"] < TOL_RATE and worst["retention_usd"] < TOL_USD and worst["fair_value_usd"] < TOL_USD
          and worst["retention_fractions"] < TOL_RATIO and worst["retention_rates"] < TOL_RATE
          and worst["levered_irr"] < TOL_LEVERED and worst["unpaid_interest_usd"] < TOL_USD
          and worst["arrears_sweep_usd"] < TOL_USD and worst["residual_cash_usd"] < TOL_USD
          and checks["labels_match"] and checks["any_shortfall_flags_match"] and checks["reserve_draw_alloc"] < TOL_USD
          and checks["breach_months_excel"] == checks["breach_months_python"]
          and checks["pool_recon_sum_abs"] < TOL_USD and checks["cash_conservation_sum_abs"] < TOL_USD
          and checks["notes_over_cash_sum"] < TOL_USD)
    return {"scenario": scen.name, "worst_abs_diff": worst, "checks": checks, "detail": detail, "passed": bool(ok)}


def breakeven_probe(src: Path, work: Path, deal: Deal, tranche: str = "C") -> dict:
    """Set the Excel CDR just below and above the Python break-even and read the loss back."""
    be = breakeven_cdr(deal, tranche, BASE_SCENARIO)
    out = {"tranche": tranche, "python_breakeven_cdr": be}
    for label, cdr in (("below", be * 0.995), ("above", be * 1.005)):
        override = (BASE_SCENARIO.name, BASE_SCENARIO.cpr, cdr, BASE_SCENARIO.severity, BASE_SCENARIO.lag, 0)
        dst = work / f"probe_{tranche}_{label}.xlsx"
        prepare(src, dst, 1, scen_override=override)
        book = Book(src, recalc(dst, work / "recalc"))
        py_loss = float(run_waterfall(deal, replace(BASE_SCENARIO, cdr=cdr))[C.tr(tranche, C.WD)].sum())
        out[label] = {"cdr": cdr, "excel_loss": book.num(f"LOSS_{tranche}"), "python_loss": py_loss}
    out["passed"] = bool(out["below"]["excel_loss"] < 1.0 and out["above"]["excel_loss"] > 1.0
                         and abs(out["above"]["excel_loss"] - out["above"]["python_loss"]) < TOL_USD)
    return out


def main(src: Path = EXCEL_PATH, out_dir: Path = OUTPUT_DIR) -> dict:
    if find_soffice() is None:
        raise SystemExit("LibreOffice not found; cannot recalculate the workbook")
    deal = Deal()
    base_fv = fair_values(run_waterfall(deal, BASE_SCENARIO))
    work = Path(tempfile.mkdtemp(prefix="waterfall_recon_"))
    results = []
    for i, scen in enumerate(SCENARIOS, start=1):
        locked = scen.name != BASE_SCENARIO.name
        dst = work / f"scenario_{i}.xlsx"
        prepare(src, dst, i, locks=base_fv if locked else None)
        results.append(compare(Book(src, recalc(dst, work / "recalc")), deal, scen, base_fv, locked))
        print(scen.name, "PASS" if results[-1]["passed"] else "FAIL", results[-1]["worst_abs_diff"])
    # pro rata mode (latches to sequential at the first breach) on a clean, a breaching and a curing scenario
    pro_deal = replace(deal, principal_mode="pro_rata")
    for i, scen in enumerate(SCENARIOS, start=1):
        if scen.name not in ("Base", "Stress", "Rate shock +300bp"):  # the last one breaches, then cures
            continue
        dst = work / f"prorata_{i}.xlsx"
        prepare(src, dst, i, locks=base_fv if scen.name != BASE_SCENARIO.name else None, mode="pro_rata")
        r = compare(Book(src, recalc(dst, work / "recalc")), pro_deal, scen, base_fv, scen.name != BASE_SCENARIO.name)
        r["scenario"] = scen.name + " (pro_rata)"
        results.append(r)
        print(r["scenario"], "PASS" if r["passed"] else "FAIL", r["worst_abs_diff"])
    # extra runs that hit the post-retirement arrears sweep: notes written down or retired while interest is owed
    for scen in EXTRA_SCENARIOS:
        dst = work / f"extra_{scen.name.replace(' ', '_')}.xlsx"
        override = (scen.name, scen.cpr, scen.cdr, scen.severity, scen.lag, scen.rate_shift_bps)
        prepare(src, dst, 1, locks=base_fv, scen_override=override)
        r = compare(Book(src, recalc(dst, work / "recalc")), deal, scen, base_fv, True)
        results.append(r)
        print(r["scenario"], "PASS" if r["passed"] else "FAIL", r["worst_abs_diff"])
    probe = breakeven_probe(src, work, deal)
    print("break-even probe", probe)
    report = {"tolerances": {"usd": TOL_USD, "rate": TOL_RATE, "levered_irr": TOL_LEVERED, "fraction": TOL_RATIO},
              "scenarios": results, "breakeven_probe": probe,
              "all_passed": all(r["passed"] for r in results) and probe["passed"]}
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "excel_reconciliation.json").write_text(json.dumps(report, indent=2, default=float))
    shutil.rmtree(work, ignore_errors=True)
    print("ALL PASSED" if report["all_passed"] else "MISMATCH")
    return report


if __name__ == "__main__":
    main()
