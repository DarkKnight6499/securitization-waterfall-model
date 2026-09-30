"""Workbook structure tests (no recalculation needed) and the LibreOffice reconciliation."""
import pytest
import openpyxl

from waterfall.config import EXCEL_PATH, Deal, SCENARIOS
from waterfall.reconcile_excel import find_soffice, main as reconcile_main

pytestmark = pytest.mark.skipif(not EXCEL_PATH.exists(), reason="workbook not built")

TIME_SERIES_SHEETS = ("Pool", "Waterfall", "Tests", "CashFlows", "Retention_CF")
FIRST_DATA_ROW = 5
PERIOD_HEADER = "Period"


@pytest.fixture(scope="module")
def wb():
    return openpyxl.load_workbook(EXCEL_PATH)


def _value(wb, name):
    sheet, addr = wb.defined_names[name].attr_text.split("!")
    return wb[sheet][addr.replace("$", "")].value


def test_creator_property(wb):
    assert wb.properties.creator == "Yazad Madan"


def test_expected_sheets_present(wb):
    for name in ("Inputs", "Pool", "Waterfall", "Tests", "CashFlows", "Retention", "Retention_CF", "Scenario_Summary"):
        assert name in wb.sheetnames


def test_calculation_sheets_are_formulas_not_pasted_values(wb):
    formulas = 0
    for name in TIME_SERIES_SHEETS:
        ws = wb[name]
        headers = {c.column: c.value for c in ws[4]}
        for row in ws.iter_rows(min_row=FIRST_DATA_ROW):
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    formulas += 1
                elif isinstance(c.value, (int, float)) and headers.get(c.column) != PERIOD_HEADER:
                    assert c.value in (0, 1), f"pasted number in {name}!{c.coordinate}"  # zero seeds and the DF of 1 at t=0 only
    assert formulas > 20000


def test_inputs_match_python_config(wb):
    deal = Deal()
    assert _value(wb, "IN_BAL_A") == deal.bal_a
    assert _value(wb, "IN_BAL_B") == deal.bal_b
    assert _value(wb, "IN_BAL_C") == deal.bal_c
    assert _value(wb, "IN_MARGIN_A") == deal.margin_a
    assert _value(wb, "IN_MARGIN_B") == deal.margin_b
    assert _value(wb, "IN_FIXED_C") == deal.fixed_c
    assert _value(wb, "IN_SVC") == deal.servicing_fee
    assert _value(wb, "IN_TRUSTEE") == deal.trustee_fee
    assert _value(wb, "IN_RES_INIT") == deal.reserve_initial
    assert _value(wb, "IN_RES_TARGET") == deal.reserve_target
    assert _value(wb, "IN_FWD_START") == deal.fwd_start
    assert _value(wb, "IN_FWD_END") == deal.fwd_end
    assert _value(wb, "IN_FWD_TAU") == deal.fwd_tau
    assert (_value(wb, "IN_OC_A"), _value(wb, "IN_OC_B"), _value(wb, "IN_OC_C")) == deal.oc_trig
    assert (_value(wb, "IN_IC_A"), _value(wb, "IN_IC_B"), _value(wb, "IN_IC_C")) == deal.ic_trig
    assert _value(wb, "IN_RET_PCT") == deal.retention_pct
    assert _value(wb, "IN_LSHAPE_V") == deal.lshape_vertical
    assert _value(wb, "IN_FIN_SPREAD") == deal.financing_spread
    assert _value(wb, "IN_MODE") == deal.principal_mode


def test_rep_lines_and_scenario_table_match_config(wb):
    ws = wb["Inputs"]
    deal = Deal()
    rows = {r[0].value: r for r in ws.iter_rows(min_col=1, max_col=6) if r[0].value}
    for name, bal, wac, wam in deal.rep_lines:
        assert (rows[name][1].value, rows[name][2].value, rows[name][3].value) == (bal, wac, wam)
    for s in SCENARIOS:
        r = rows[s.name]
        assert (r[1].value, r[2].value, r[3].value, r[4].value, r[5].value) == (
            s.cpr, s.cdr, s.severity, s.lag, s.rate_shift_bps)


@pytest.mark.skipif(find_soffice() is None, reason="LibreOffice not available")
def test_excel_recalculation_reconciles_to_python():
    report = reconcile_main(EXCEL_PATH)
    assert report["all_passed"], report


def _numeric_literals(ws):
    return {c.coordinate: c.value for row in ws.iter_rows() for c in row if isinstance(c.value, (int, float))}


def test_summary_sheets_hold_only_documented_structural_zeros(wb):
    # R2 held in non-L options, R2 par in non-carve-out options, and the R row of the note-only columns
    assert _numeric_literals(wb["Retention"]) == {c: 0 for c in ("F15", "G15", "F16", "F18", "G18", "F19", "G19")}
    assert _numeric_literals(wb["Scenario_Summary"]) == {c: 0 for c in ("F10", "M10", "N10")}
